"""Planos padrao (Free, Basico, Premium, Premium + Estacao), plano gratis
liberado na hora e preco/abas editaveis pelo Master.

Rodar a partir de webapp/:  python -m unittest tests.test_planos
"""
import os
import sys
import tempfile
import unittest

os.environ.setdefault("BIOSCOUT_DB_PATH", os.path.join(tempfile.mkdtemp(), "teste.db"))
os.environ.pop("ASAAS_API_KEY", None)
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import app as webapp  # noqa: E402
import models  # noqa: E402

CPF_VALIDO = "52998224725"


class PlanosTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        webapp.app.config["TESTING"] = True
        models.init_db()
        if not models.get_user_by_username("Zzz Teste Master Planos"):
            models.create_user("Zzz Teste Master Planos", "x", is_master=True)
        cls.planos = {p["nome"]: p for p in models.get_planos()}

    def _logado(self, user_id):
        c = webapp.app.test_client()
        with c.session_transaction() as s:
            s["_user_id"] = str(user_id)
            s["_fresh"] = True
        return c

    def _cliente(self, nome):
        return models.criar_cliente(nome, "segredo1", "", "5565999990000", "", None)

    def test_planos_padrao(self):
        for nome in ("Free", "Básico", "Premium", "Premium + Estação"):
            self.assertIn(nome, self.planos)
        free = self.planos["Free"]
        self.assertEqual((free["valor"], free["ativo"], free["recursos"]), (0, 1, ["clima"]))
        self.assertEqual(self.planos["Premium"]["ativo"], 0)  # pago nasce fora de venda ate' ter preco
        models.init_db()  # rodar de novo nao duplica
        self.assertEqual(sum(1 for p in models.get_planos() if p["nome"] == "Free"), 1)

    def test_master_define_preco_e_abas(self):
        master = self._logado(models.get_user_by_username("Zzz Teste Master Planos")["id"])
        premium = self.planos["Premium"]
        basico = self.planos["Básico"]
        k, b = premium["id"], basico["id"]
        # um envio so' grava todos os planos da tela
        master.post("/admin/planos/salvar", data={
            "plano_ids": [str(k), str(b)],
            f"nome__{k}": "Premium", f"valor__{k}": "149,90", f"ciclo__{k}": "MONTHLY", f"max_pontos__{k}": "10",
            f"raio_km__{k}": "50", f"ordem__{k}": "3", f"recursos__{k}": ["clima", "esporos", "risco"], f"ativo__{k}": "on",
            f"nome__{b}": "Básico", f"valor__{b}": "59,90", f"ciclo__{b}": "MONTHLY", f"max_pontos__{b}": "3",
            f"raio_km__{b}": "50", f"ordem__{b}": "2", f"recursos__{b}": ["clima", "site_mapa"],
            "nome__novo": "",
        })
        p = models.get_plano(k)
        self.assertEqual((p["valor"], p["ativo"], p["recursos"]), (149.9, 1, ["clima", "esporos", "risco"]))
        self.assertEqual((models.get_plano(b)["valor"], models.get_plano(b)["ativo"]), (59.9, 0))
        html = master.get("/admin/planos").get_data(as_text=True)
        self.assertIn("Abas liberadas no app", html)
        self.assertIn("0 = grátis", html)
        self.assertEqual(html.count('type="submit">Salvar</button>'), 1)  # um botao Salvar so'
        # com um plano errado, nenhum e' gravado
        negativo = master.post("/admin/planos/salvar", data={
            "plano_ids": [str(k), str(b)],
            f"nome__{k}": "Premium", f"valor__{k}": "-1", f"ciclo__{k}": "MONTHLY", f"max_pontos__{k}": "10", f"raio_km__{k}": "50",
            f"nome__{b}": "Básico", f"valor__{b}": "1,00", f"ciclo__{b}": "MONTHLY", f"max_pontos__{b}": "3", f"raio_km__{b}": "50",
        }, follow_redirects=True)
        self.assertIn("0 = plano grátis", negativo.get_data(as_text=True))
        self.assertEqual((models.get_plano(k)["valor"], models.get_plano(b)["valor"]), (149.9, 59.9))

    def test_free_libera_na_hora_sem_cpf(self):
        uid = self._cliente("Zzz Cliente Free")
        c = self._logado(uid)
        html = c.get("/app/assinatura").get_data(as_text=True)
        self.assertIn("Escolha o seu plano", html)
        self.assertIn("Grátis", html)
        # todos os planos aparecem; os fora de venda como "Em breve", sem escolha
        self.assertIn("Premium + Estação", html)
        self.assertIn("Em breve", html)
        self.assertIn(f'id="plano_{self.planos["Premium + Estação"]["id"]}"', html)
        self.assertNotIn("Aguardando pagamento", html)  # situacao so' depois de escolher
        forcado = c.post("/app/assinatura/assinar", data={"plano_id": str(self.planos["Premium + Estação"]["id"])},
                         follow_redirects=True)
        self.assertIn("Escolha um plano", forcado.get_data(as_text=True))
        self.assertIsNone(models.get_assinatura(uid)["plano_id"])
        self.assertIn("/app/assinatura", c.get("/app").headers["Location"])  # sem plano: bloqueado
        resp = c.post("/app/assinatura/assinar", data={"plano_id": str(self.planos["Free"]["id"])})
        self.assertTrue(resp.headers["Location"].endswith("/app"))
        assinatura = models.get_assinatura(uid)
        self.assertEqual((assinatura["plano_id"], assinatura["status"], assinatura["cpf_cnpj"]),
                         (self.planos["Free"]["id"], "ativa", ""))
        self.assertTrue(models.assinatura_liberada(assinatura))
        self.assertEqual(c.get("/app").status_code, 200)
        self.assertIn("Sem prazo", c.get("/app/assinatura").get_data(as_text=True))

    def test_trocar_free_por_pago(self):
        basico = self.planos["Básico"]
        models.salvar_plano(basico["id"], "Básico", "", 49.9, "MONTHLY", 3, 50, ["clima", "esporos"], True, 2)
        uid = self._cliente("Zzz Cliente Troca")
        c = self._logado(uid)
        c.post("/app/assinatura/assinar", data={"plano_id": str(self.planos["Free"]["id"])})
        # sem CPF: recusa e continua no gratis
        html = c.post("/app/assinatura/assinar", data={"plano_id": str(basico["id"])}, follow_redirects=True).get_data(as_text=True)
        self.assertIn("CPF ou CNPJ válido", html)
        self.assertTrue(models.assinatura_liberada(models.get_assinatura(uid)))
        # com CPF: vira o pago, aguardando pagamento (Asaas nao configurado no teste)
        c.post("/app/assinatura/assinar", data={"plano_id": str(basico["id"]), "cpf_cnpj": CPF_VALIDO})
        assinatura = models.get_assinatura(uid)
        self.assertEqual((assinatura["plano_id"], assinatura["cpf_cnpj"], assinatura["cortesia_ate"]),
                         (basico["id"], CPF_VALIDO, None))
        self.assertFalse(models.assinatura_liberada(assinatura))


    def _cliente_no_plano(self, nome, recursos):
        plano_id = models.salvar_plano(None, f"Zzz Plano {nome}", "", 0, "MONTHLY", 3, 40, recursos, True, 9)
        uid = self._cliente(nome)
        models.atualizar_assinatura(uid, plano_id=plano_id, status="ativa", cortesia_ate=models.SEM_PRAZO)
        return uid

    def test_recursos_do_site_depois_da_previsao(self):
        chaves = [c for c, _ in models.RECURSOS_PLANO]
        self.assertEqual(chaves[:6], ["clima", "site_mapa", "site_alertas", "site_graficos", "site_manejos", "site_recomendacoes"])
        self.assertIn("fazenda_estimada", chaves)

    def test_abas_do_site_conforme_plano_e_fazenda_estimada(self):
        uid = self._cliente_no_plano("Abas", ["clima", "site_mapa", "site_manejos", "site_recomendacoes", "fazenda_estimada"])
        c = self._logado(uid)
        self.assertEqual(c.get("/mapa").status_code, 200)
        self.assertEqual(c.get("/fazendas").status_code, 200)
        self.assertEqual(c.get("/recommendations", follow_redirects=True).status_code, 200)
        for fora in ("/graficos", "/dashboard", "/admin/users", "/mapa-interpolado"):
            with self.subTest(rota=fora):
                self.assertTrue(c.get(fora).headers["Location"].endswith("/app"))
        menu = c.get("/mapa").get_data(as_text=True)
        self.assertIn("📱 App", menu)
        self.assertNotIn(">Graficos<", menu)
        self.assertIn("🌐", c.get("/app").get_data(as_text=True))

        # lavoura nova vira fazenda estimada liberada so' pra ele
        resp = c.post("/app/pontos/novo", data={"nome": "Talhao Zzz", "lat": "-12.5", "lon": "-55.7"})
        self.assertEqual(resp.status_code, 302)
        ponto = models.get_cliente_pontos(uid)[0]
        self.assertTrue(ponto["site_name"])
        vf = models.get_virtual_farm(ponto["site_name"])
        self.assertEqual((vf["nome"], vf["raio_km"], vf["tipo"]), ("Talhao Zzz (Abas)", 40, "doenca"))
        self.assertIn(ponto["site_name"], models.get_user_permitted_site_names(uid))
        # apagar a lavoura apaga a fazenda estimada
        c.post(f"/app/pontos/{ponto['id']}/excluir")
        self.assertIsNone(models.get_virtual_farm(ponto["site_name"]))

    def test_plano_sem_fazenda_estimada_e_free_sem_site(self):
        uid = self._cliente_no_plano("Estacao", ["clima", "site_mapa"])
        c = self._logado(uid)
        c.post("/app/pontos/novo", data={"nome": "Talhao Estacao", "lat": "-12.5", "lon": "-55.7"})
        self.assertIsNone(models.get_cliente_pontos(uid)[0]["site_name"])  # Premium + Estacao: Master libera a real
        self.assertEqual(c.get("/mapa").status_code, 200)
        free = self._cliente_no_plano("Free Site", ["clima"])
        self.assertTrue(self._logado(free).get("/mapa").headers["Location"].endswith("/app"))

    def test_excluir_cliente_apaga_fazendas_estimadas(self):
        uid = self._cliente_no_plano("Excluir Site", ["clima", "site_mapa", "fazenda_estimada"])
        self._logado(uid).post("/app/pontos/novo", data={"nome": "Talhao Exc", "lat": "-12.5", "lon": "-55.7"})
        site_name = models.get_cliente_pontos(uid)[0]["site_name"]
        self.assertIsNotNone(models.get_virtual_farm(site_name))
        models.delete_user(uid)
        self.assertIsNone(models.get_virtual_farm(site_name))


class ExcelClientesPlanosTest(unittest.TestCase):
    """Abas "Clientes" e "Planos" no Excel de Exportar (pedido do Alan, 01/11/2026)."""

    def test_abas_clientes_e_planos(self):
        import datetime as dt
        import export_excel
        from openpyxl import load_workbook
        plano = models.salvar_plano(None, "Zzz Plano Excel", "", 79.9, "MONTHLY", 3, 50, ["clima", "esporos"], True, 9)
        uid = models.criar_cliente("Zzz Cliente Excel", "segredo1", "excel@t.com", "5565999990055", "", plano)
        models.atualizar_assinatura(uid, plano_id=plano, status="ativa", cortesia_ate=models.SEM_PRAZO)
        models.registrar_cadastro({"username": "Zzz Cliente Excel", "nome_completo": "Zzz Cliente Excel da Silva",
                                   "perfil": "produtor,tecnico", "telefone": "5565999990055", "empresa": "Fazenda Zzz",
                                   "cidade": "Sorriso", "uf": "MT", "culturas": "Soja", "area_ha": 1500,
                                   "tipo": "cliente", "user_id": uid})
        wb = load_workbook(export_excel.build_workbook())
        self.assertIn("Clientes", wb.sheetnames)
        self.assertIn("Planos", wb.sheetnames)
        clientes = list(wb["Clientes"].iter_rows(values_only=True))
        self.assertEqual(clientes[0][:3], ("Cliente", "Usuário", "Celular"))
        linha = next(r for r in clientes if r[1] == "Zzz Cliente Excel")
        cab = clientes[0]
        valor = dict(zip(cab, linha))
        self.assertEqual(valor["Cliente"], "Zzz Cliente Excel da Silva")
        self.assertEqual(valor["Perfil"], "Produtor, Técnico do setor")
        self.assertEqual((valor["Cidade"], valor["Plano"], valor["Valor do plano (R$)"], valor["Ciclo"]),
                         ("Sorriso/MT", "Zzz Plano Excel", 79.9, "Mensal"))
        self.assertEqual((valor["Situação"], valor["Acesso liberado"], valor["Cortesia até"]), ("Ativa", "Sim", "Sem prazo"))
        self.assertIsInstance(valor["Cadastrado em"], dt.datetime)  # data de verdade, nao texto
        self.assertNotIn("Erro ao gerar", str(clientes))
        planos = list(wb["Planos"].iter_rows(values_only=True))
        p = dict(zip(planos[0], next(r for r in planos if r[0] == "Zzz Plano Excel")))
        self.assertEqual((p["Valor (R$)"], p["À venda"], p["Clientes no plano"], p["Clientes com acesso liberado"]),
                         (79.9, "Sim", 1, 1))
        self.assertIn("Esporos", p["Abas liberadas"])
        self.assertIn("R$", wb["Planos"]["B2"].number_format)


if __name__ == "__main__":
    unittest.main()
